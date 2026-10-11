"""Import bank statements / ledger exports (CSV, XLSX) into normalised transactions.

Sign convention for normalised `amount`: positive = money INTO the bank account.
 - bank statement: credit => +, debit => -
 - cash book (company's bank ledger): debit (receipt) => +, credit (payment) => -
"""
import csv
import io
import re
from datetime import date, datetime

HEADER_SYNONYMS = {
    "date": ["date", "txn date", "transaction date", "value date", "posting date", "tran date"],
    "description": ["description", "narration", "particulars", "details", "remarks", "transaction remarks"],
    "reference": ["reference", "ref no", "ref", "cheque no", "chq no", "chq/ref no", "utr", "voucher no", "instrument no", "cheque/ref no"],
    "debit": ["debit", "withdrawal", "withdrawals", "dr", "withdrawal amt", "debit amount", "paid out"],
    "credit": ["credit", "deposit", "deposits", "cr", "deposit amt", "credit amount", "paid in"],
    "amount": ["amount", "txn amount", "transaction amount"],
    "balance": ["balance", "closing balance", "running balance"],
}
_DATE_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d %b %Y", "%d-%b-%Y", "%d/%m/%y", "%d-%b-%y"]


class ParseError(ValueError):
    pass


def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _to_date(v) -> str | None:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    for f in _DATE_FORMATS:
        try:
            return datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    return None


def _to_num(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("₹", "").strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    s = re.sub(r"(?i)\s*(dr|cr)\.?$", "", s)
    try:
        n = float(s)
    except ValueError:
        return 0.0
    return -n if neg else n


def _read_rows(filename: str, data: bytes) -> list[list]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    return list(csv.reader(io.StringIO(text)))


def _map_columns(header: list) -> dict[str, int]:
    cols: dict[str, int] = {}
    for idx, h in enumerate(header):
        nh = _norm_header(h)
        for field, syns in HEADER_SYNONYMS.items():
            if field not in cols and nh in syns:
                cols[field] = idx
    return cols


def parse_statement(filename: str, data: bytes, side: str, prefix: str) -> list[dict]:
    """side: 'bank' or 'book'. prefix is used to build stable txn ids (e.g. 'B', 'L')."""
    rows = _read_rows(filename, data)
    # Statements often have preamble lines: find the first row that looks like a header.
    for hi, row in enumerate(rows):
        cols = _map_columns(row)
        if "date" in cols and ("amount" in cols or "debit" in cols or "credit" in cols):
            break
    else:
        raise ParseError("Could not find a header row with a date and amount/debit/credit columns")

    txns = []
    for row in rows[hi + 1:]:
        if not row or len(row) <= cols["date"]:
            continue
        d = _to_date(row[cols["date"]])
        if d is None:
            continue  # totals, blank lines, footers
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None
        if "amount" in cols and "debit" not in cols and "credit" not in cols:
            amt = _to_num(get("amount"))
        else:
            cr, dr = _to_num(get("credit")), _to_num(get("debit"))
            amt = (cr - dr) if side == "bank" else (dr - cr)
        if amt == 0:
            continue
        txns.append({
            "id": f"{prefix}{len(txns) + 1}", "date": d,
            "description": str(get("description") or "").strip(),
            "reference": str(get("reference") or "").strip(),
            "amount": round(amt, 2),
        })
    if not txns:
        raise ParseError("No transactions found")
    return txns
