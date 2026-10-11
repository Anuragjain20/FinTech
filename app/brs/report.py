import io

from openpyxl import Workbook
from openpyxl.styles import Font


def build_xlsx(rec: dict, summary: dict, classified: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws.append(["Bank Reconciliation Statement"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in summary.items():
        ws.append([k.replace("_", " ").capitalize(), v])
    ws.column_dimensions["A"].width = 36

    def sheet(title, headers, rows):
        s = wb.create_sheet(title)
        s.append(headers)
        for c in s[1]:
            c.font = bold
        for r in rows:
            s.append(r)

    bank = {t["id"]: t for t in rec["bank"]}
    book = {t["id"]: t for t in rec["book"]}
    sheet("Matched", ["Bank date", "Bank desc", "Book date", "Book desc", "Amount", "Rule", "Confidence", "Status"],
          [[bank[m["bank_id"]]["date"], bank[m["bank_id"]]["description"], book[m["book_id"]]["date"],
            book[m["book_id"]]["description"], bank[m["bank_id"]]["amount"], m["rule"], m["confidence"], m["status"]]
           for m in rec["matches"]])
    cols = ["id", "date", "description", "reference", "amount", "category", "explanation"]
    sheet("Unmatched - Bank", cols, [[t[c] for c in cols] for t in classified["bank"]])
    sheet("Unmatched - Books", cols, [[t[c] for c in cols] for t in classified["book"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
