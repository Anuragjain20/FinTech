import io

from docx import Document


def to_docx(title: str, text: str) -> bytes:
    d = Document()
    d.add_heading(title, level=1)
    d.add_paragraph().add_run("DRAFT - for review by a qualified legal professional before signing.").italic = True
    for para in text.split("\n\n")[1:]:
        lines = para.strip().split("\n")
        if len(lines) > 1 and lines[0][:1].isdigit():
            d.add_heading(lines[0], level=2)
            d.add_paragraph("\n".join(lines[1:]))
        else:
            d.add_paragraph(para.strip())
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
