import io

from openpyxl import load_workbook

BANK = """HDFC Bank Statement
Account: 1234
Date,Narration,Chq/Ref No,Withdrawal Amt,Deposit Amt,Closing Balance
01/04/2024,NEFT-ACME-INV101,UTR1001,,50000.00,150000
02/04/2024,CHQ PAID 000123,000123,20000.00,,130000
05/04/2024,SMS CHARGES,,25.00,,129975
07/04/2024,CUSTOMER PAYMENT RAMESH,,,7500.00,137475
30/04/2024,INTEREST PD,,,310.00,137785
"""
BOOK = """Date,Particulars,Voucher No,Debit,Credit
01/04/2024,Receipt from Acme inv 101,UTR1001,50000,
03/04/2024,Cheque to supplier,123,,20000
06/04/2024,Receipt Ramesh,,7500,
29/04/2024,Cheque to landlord,000999,,5000
"""
# books closing = 100000 + 50000 - 20000 + 7500 - 5000 ; bank closing per statement
BOOK_CLOSING = 132500.0
BANK_CLOSING = 137785.0 - 100000.0 + 100000.0


def create(client, headers, **kw):
    data = {"name": "Apr", "bank_closing_balance": str(BANK_CLOSING), "book_closing_balance": str(BOOK_CLOSING)}
    data.update(kw)
    return client.post("/reconciliations", headers=headers, data=data,
                       files={"bank_file": ("bank.csv", BANK.encode()), "book_file": ("book.csv", BOOK.encode())})


def test_matching_rules_and_classification(client, admin):
    r = create(client, admin)
    assert r.status_code == 201, r.text
    v = r.json()
    by_rule = {m["bank_id"]: m for m in v["matches"]}
    assert by_rule["B1"]["rule"] == "reference" and by_rule["B1"]["status"] == "matched"
    assert by_rule["B2"]["rule"] == "reference"  # 000123 vs 123: leading zeros normalised
    assert by_rule["B4"]["rule"] == "amount+date"
    cats = {t["id"]: t["category"] for t in v["unmatched_bank"]}
    assert cats == {"B3": "bank_charge", "B5": "interest"}
    assert [t["id"] for t in v["unmatched_book"]] == ["L4"]
    s = v["summary"]
    # bank-only: -25 + 310 = 285 ; book-only: -5000 ; expected bank = 132500 + 285 + 5000
    assert s["bank_only_net"] == 285.0 and s["book_only_net"] == -5000.0
    assert s["unexplained_difference"] == round(BANK_CLOSING - (BOOK_CLOSING + 285 + 5000), 2)


def test_perfect_reconciliation_and_finalize(client, admin, role_headers):
    bank_close = BOOK_CLOSING + 285 + 5000
    v = create(client, admin, bank_closing_balance=str(bank_close)).json()
    assert v["summary"]["unexplained_difference"] == 0 and v["summary"]["reconciled"] is True
    rid = v["id"]
    assert client.post(f"/reconciliations/{rid}/finalize", headers=role_headers("operator")).status_code == 403
    assert client.post(f"/reconciliations/{rid}/finalize", headers=admin).json()["status"] == "finalized"
    assert client.post(f"/reconciliations/{rid}/manual-match", json={"bank_id": "B3", "book_id": "L4"}, headers=admin).status_code == 409


def test_fuzzy_suggestion_must_be_confirmed(client, admin):
    bank = "Date,Description,Debit,Credit\n10/04/2024,PAYMENT VENDOR XYZ ENTERPRISES,900,\n"
    book = "Date,Description,Debit,Credit\n17/04/2024,Vendor XYZ Enterprises payment,,900\n"
    r = client.post("/reconciliations", headers=admin, data={"book_closing_balance": "0", "bank_closing_balance": "0"},
                    files={"bank_file": ("b.csv", bank.encode()), "book_file": ("l.csv", book.encode())}).json()
    m = r["matches"][0]
    assert m["rule"] == "amount+fuzzy" and m["status"] == "suggested"
    rid = r["id"]
    assert client.post(f"/reconciliations/{rid}/finalize", headers=admin).status_code == 409
    done = client.post(f"/reconciliations/{rid}/confirm", json={"bank_id": "B1", "book_id": "L1"}, headers=admin).json()
    assert done["matches"][0]["status"] == "confirmed"
    assert client.post(f"/reconciliations/{rid}/finalize", headers=admin).status_code == 200


def test_reject_then_manual_and_audit(client, admin):
    bank = "Date,Description,Debit,Credit\n10/04/2024,PAYMENT VENDOR XYZ,900,\n"
    book = "Date,Description,Debit,Credit\n17/04/2024,Vendor XYZ payment,,900\n"
    rid = client.post("/reconciliations", headers=admin,
                      files={"bank_file": ("b.csv", bank.encode()), "book_file": ("l.csv", book.encode())}).json()["id"]
    v = client.post(f"/reconciliations/{rid}/reject", json={"bank_id": "B1", "book_id": "L1"}, headers=admin).json()
    assert v["matches"] == [] and len(v["unmatched_bank"]) == 1
    v = client.post(f"/reconciliations/{rid}/manual-match", json={"bank_id": "B1", "book_id": "L1"}, headers=admin).json()
    assert v["matches"][0]["rule"] == "manual"
    actions = [a["action"] for a in client.get(f"/reconciliations/{rid}", headers=admin).json()["audit"]]
    assert actions == ["reconciliation.created", "match.rejected", "match.manual"]


def test_xlsx_input_and_report(client, admin):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Date", "Narration", "Debit", "Credit"]); ws.append(["01/04/2024", "x", None, 100])
    buf = io.BytesIO(); wb.save(buf)
    book = "Date,Description,Debit,Credit\n01/04/2024,x,100,\n"
    r = client.post("/reconciliations", headers=admin, files={"bank_file": ("b.xlsx", buf.getvalue()), "book_file": ("l.csv", book.encode())})
    assert r.status_code == 201 and r.json()["matches"][0]["rule"] == "amount+date"
    rep = client.get(f"/reconciliations/{r.json()['id']}/report.xlsx", headers=admin)
    assert rep.status_code == 200
    assert load_workbook(io.BytesIO(rep.content)).sheetnames == ["Summary", "Matched", "Unmatched - Bank", "Unmatched - Books"]


def test_bad_file_422_and_isolation(client, admin, other_admin):
    r = client.post("/reconciliations", headers=admin, files={"bank_file": ("b.csv", b"foo,bar\n1,2"), "book_file": ("l.csv", b"a")})
    assert r.status_code == 422
    rid = create(client, admin).json()["id"]
    assert client.get(f"/reconciliations/{rid}", headers=other_admin).status_code == 404
